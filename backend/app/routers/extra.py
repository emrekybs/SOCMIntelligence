"""
Ek uc noktalar: ayarlar, avatar hash (kimlik eslestirme), delil dogrulama, Excel disa aktarma, izleme.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import re
from datetime import datetime, timezone
from typing import Any

import requests
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from app.core import settings_store, watch
from app.core.cache import cache
from app.tasks.scan_tasks import ADAPTERS, _get_adapter, sha256_of

router = APIRouter(prefix="/api", tags=["extra"])


# ── Ayarlar ────────────────────────────────────────────────────────
@router.get("/settings")
def get_settings() -> dict[str, Any]:
    return settings_store.public_view()


@router.put("/settings")
def put_settings(patch: dict[str, Any]) -> dict[str, Any]:
    try:
        settings_store.update(patch)
    except (ValueError, TypeError) as e:
        raise HTTPException(status_code=400, detail=str(e))
    return settings_store.public_view()


@router.post("/settings/test-proxy")
def test_proxy() -> dict[str, Any]:
    """Uygulanan proxy uzerinden dis IP'yi ve Tor kullanimini kontrol et."""
    try:
        r = requests.get("https://check.torproject.org/api/ip", timeout=20)
        j = r.json()
        return {"ok": True, "ip": j.get("IP"), "tor": bool(j.get("IsTor")), "proxy": settings_store.proxy_url(settings_store.load())}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": str(e)[:300], "proxy": settings_store.proxy_url(settings_store.load())}


# ── Avatar hash (ayni kisi skoru) ──────────────────────────────────
@router.get("/socmint/avatar-hash")
async def avatar_hash(url: str = Query(..., min_length=8, max_length=2000)) -> dict[str, Any]:
    if not re.match(r"^https?://", url):
        raise HTTPException(status_code=400, detail="Gecersiz URL")
    key = "avatar:" + hashlib.sha1(url.encode()).hexdigest()
    cached = await cache.get(key)
    if cached:
        return cached

    def work():
        import imagehash
        from PIL import Image
        r = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
        r.raise_for_status()
        if len(r.content) > 8_000_000:
            raise ValueError("Görsel çok büyük")
        img = Image.open(io.BytesIO(r.content)).convert("RGB")
        return {"url": url, "phash": str(imagehash.phash(img)), "dhash": str(imagehash.dhash(img)),
                "md5": hashlib.md5(r.content).hexdigest(), "size": img.size}
    try:
        res = await asyncio.to_thread(work)
    except Exception as e:  # noqa: BLE001
        return {"url": url, "phash": None, "error": f"Avatar alınamadı: {str(e)[:200]}"}
    await cache.set(key, res, ttl=86400)
    return res


# ── Delil dogrulama ─────────────────────────────────────────────────
class VerifyReq(BaseModel):
    data: Any
    sha256: str = Field(..., min_length=64, max_length=64)


@router.post("/evidence/verify")
def verify(req: VerifyReq) -> dict[str, Any]:
    actual = sha256_of(req.data)
    return {"match": actual == req.sha256.lower(), "expected": req.sha256.lower(), "actual": actual}


# ── Excel ────────────────────────────────────────────────────────────
class XlsxReq(BaseModel):
    title: str = "SOCMIntelligence"
    scans: list[dict[str, Any]] = []
    graph: dict[str, Any] = {}
    annotations: dict[str, Any] = {}
    lang: str = "tr"


# Excel sayfa adlari ve sutun basliklari (req.lang ile secilir; varsayilan tr)
_XLSX_LABELS: dict[str, dict[str, Any]] = {
    "tr": {
        "summary": "Özet", "entities": "Varlıklar", "relations": "İlişkiler",
        "summary_cols": ["#", "Platform", "Girdi", "Değer", "Durum", "Tarama zamanı", "SHA-256"],
        "entity_cols": ["ID", "Tür", "Değer", "Kaynaklar", "Analist işareti", "Analist notu", "Meta"],
        "relation_cols": ["Kaynak", "Kaynak ID", "İlişki", "Hedef", "Hedef ID", "Ağırlık"],
        "data_cols": ["Alan", "Değer"],
    },
    "en": {
        "summary": "Summary", "entities": "Entities", "relations": "Relations",
        "summary_cols": ["#", "Platform", "Input", "Value", "Status", "Scan time", "SHA-256"],
        "entity_cols": ["ID", "Type", "Value", "Sources", "Analyst mark", "Analyst note", "Meta"],
        "relation_cols": ["Source", "Source ID", "Relation", "Target", "Target ID", "Weight"],
        "data_cols": ["Field", "Value"],
    },
}


def _flat_rows(obj: Any, prefix: str = "", out: list | None = None, depth: int = 0) -> list:
    out = [] if out is None else out
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flat_rows(v, f"{prefix}.{k}" if prefix else str(k), out, depth + 1)
    elif isinstance(obj, list):
        if obj and all(not isinstance(x, (dict, list)) for x in obj):
            out.append((prefix, ", ".join(str(x) for x in obj)))
        else:
            for i, v in enumerate(obj[:500]):
                _flat_rows(v, f"{prefix}[{i}]", out, depth + 1)
    else:
        out.append((prefix, obj))
    return out


@router.post("/export/xlsx")
def export_xlsx(req: XlsxReq):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3D6EA5")

    def sheet(ws, cols, rows, widths=None):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append([("" if v is None else (v if isinstance(v, (int, float)) else str(v)[:32000])) for v in r])
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    L = _XLSX_LABELS.get((req.lang or "tr").lower(), _XLSX_LABELS["tr"])
    ws = wb.active
    ws.title = L["summary"]
    sheet(ws, L["summary_cols"],
          [[i + 1, s.get("platform"), s.get("raw"), s.get("value"), s.get("status"),
            datetime.fromtimestamp((s.get("at") or 0) / 1000, tz=timezone.utc).isoformat(timespec="seconds") if s.get("at") else "",
            (s.get("evidence") or {}).get("sha256")] for i, s in enumerate(req.scans)], [5, 14, 30, 30, 12, 24, 70])
    ann = req.annotations or {}
    sheet(wb.create_sheet(L["entities"]), L["entity_cols"],
          [[n.get("id"), n.get("type"), n.get("label"), ", ".join(n.get("sources") or []),
            (ann.get(n.get("id")) or {}).get("status"), (ann.get(n.get("id")) or {}).get("note"),
            "; ".join(f"{k}={v}" for k, v in (n.get("meta") or {}).items())] for n in (req.graph.get("nodes") or [])], [40, 12, 34, 30, 14, 30, 80])
    labels = {n.get("id"): n.get("label") for n in (req.graph.get("nodes") or [])}
    sheet(wb.create_sheet(L["relations"]), L["relation_cols"],
          [[labels.get(e.get("source")), e.get("source"), e.get("type"), labels.get(e.get("target")), e.get("target"), e.get("weight")]
           for e in (req.graph.get("edges") or [])], [30, 36, 16, 30, 36, 10])
    used = set()
    for i, s in enumerate(req.scans):
        if not s.get("data"):
            continue
        name = re.sub(r"[\[\]\*\?/\\:]", "_", f"{i + 1}_{s.get('platform')}_{s.get('value') or ''}")[:31]
        while name in used:
            name = name[:28] + f"_{len(used)}"
        used.add(name)
        sheet(wb.create_sheet(name), L["data_cols"], _flat_rows(s["data"]), [60, 100])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = re.sub(r"[^\w.-]+", "_", req.title)[:60] or "socmintelligence"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fn}.xlsx"'})


# ── Izleme ─────────────────────────────────────────────────────────
class WatchCreate(BaseModel):
    platform: str
    raw: str = Field(..., min_length=1, max_length=300)
    interval_hours: float = Field(24, ge=1, le=24 * 30)


class WatchPatch(BaseModel):
    enabled: bool | None = None
    interval_hours: float | None = Field(None, ge=1, le=24 * 30)


@router.get("/watches")
def watches() -> dict[str, Any]:
    return {"watches": watch.list_watches()}


@router.post("/watches")
def create_watch(req: WatchCreate) -> dict[str, Any]:
    if req.platform not in ADAPTERS:
        raise HTTPException(status_code=400, detail="Bilinmeyen platform")
    try:
        t, v = _get_adapter(req.platform).classify(req.raw)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return watch.create_watch(req.platform, req.raw, t, v, req.interval_hours)


@router.patch("/watches/{wid}")
def patch_watch(wid: int, req: WatchPatch) -> dict[str, Any]:
    w = watch.update_watch(wid, enabled=req.enabled, interval_hours=req.interval_hours)
    if not w:
        raise HTTPException(status_code=404, detail="İzleme bulunamadı")
    return w


@router.delete("/watches/{wid}")
def delete_watch(wid: int) -> dict[str, Any]:
    if not watch.delete_watch(wid):
        raise HTTPException(status_code=404, detail="İzleme bulunamadı")
    return {"deleted": True}


@router.post("/watches/{wid}/run")
async def run_now(wid: int) -> dict[str, Any]:
    if not watch.get_watch(wid):
        raise HTTPException(status_code=404, detail="İzleme bulunamadı")
    w = await asyncio.to_thread(watch.run_watch, wid)
    return w


@router.get("/watches/{wid}/snapshots")
def watch_snapshots(wid: int, limit: int = Query(30, ge=1, le=200)) -> dict[str, Any]:
    return {"snapshots": watch.snapshots(wid, limit)}
